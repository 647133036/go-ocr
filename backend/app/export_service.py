"""导出服务：将 OCR 结构化结果导出为 Word / Excel / PDF。"""

import html
import io
import re

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook


def parse_html_table(html_str: str) -> list[list[str]]:
    """将 <table> HTML 字符串解析为二维列表。"""
    rows = re.findall(r"<tr[^>]*>(.*?)</tr>", html_str, re.S | re.I)
    table: list[list[str]] = []
    for row in rows:
        cells = re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", row, re.S | re.I)
        cleaned = []
        for cell in cells:
            text = re.sub(r"<[^>]+>", "", cell)
            text = html.unescape(text).strip()
            cleaned.append(text)
        if cleaned:
            table.append(cleaned)
    return table


def _iter_blocks(result: dict):
    for page in result.get("pages", []):
        for block in page.get("blocks", []):
            yield block


def export_docx(result: dict) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)

    for block in _iter_blocks(result):
        btype = block.get("type")
        if btype in ("text", "formula"):
            doc.add_paragraph(block.get("text", ""))
        elif btype == "table":
            table_data = parse_html_table(block.get("html", ""))
            if not table_data:
                continue
            rows = len(table_data)
            cols = max(len(r) for r in table_data)
            table = doc.add_table(rows=rows, cols=cols)
            table.style = "Table Grid"
            for i, row in enumerate(table_data):
                for j in range(cols):
                    table.rows[i].cells[j].text = row[j] if j < len(row) else ""

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_xlsx(result: dict) -> bytes:
    wb = Workbook()
    tables: list[list[list[str]]] = []
    texts: list[str] = []

    for block in _iter_blocks(result):
        btype = block.get("type")
        if btype == "table":
            tables.append(parse_html_table(block.get("html", "")))
        elif btype in ("text", "formula"):
            texts.append(block.get("text", ""))

    if tables:
        for idx, table in enumerate(tables, 1):
            ws = wb.active if idx == 1 else wb.create_sheet()
            ws.title = f"表格{idx}"
            for row in table:
                ws.append(row)
    else:
        ws = wb.active
        ws.title = "文本"
        ws.append(["段落"])
        for text in texts:
            ws.append([text])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_pdf(result: dict) -> bytes:
    """生成双层 PDF：原图作为底层，叠加不可见文字层以保留版式。"""
    import fitz

    doc = fitz.open()
    for page in result.get("pages", []):
        image_path = page.get("image_path")
        if not image_path:
            continue

        img_doc = fitz.open(image_path)
        rect = img_doc[0].rect
        pdfpage = doc.new_page(width=rect.width, height=rect.height)
        pdfpage.insert_image(pdfpage.rect, filename=image_path)
        img_doc.close()

        for block in page.get("blocks", []):
            if block.get("type") not in ("text", "formula"):
                continue
            text = block.get("text", "")
            bbox = block.get("bbox")
            if not text or not bbox or len(bbox) != 4:
                continue
            x1, y1, x2, y2 = bbox
            fontsize = max(6.0, min(14.0, (y2 - y1) * 0.8))
            r = fitz.Rect(x1, y1, x2, y2)
            try:
                pdfpage.insert_textbox(
                    r,
                    text,
                    fontsize=fontsize,
                    fontname="china-s",
                    render_mode=3,
                )
            except Exception:
                continue

    buf = doc.tobytes()
    doc.close()
    return buf
